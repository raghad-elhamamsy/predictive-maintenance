"""Read a generic workbook of timestamped, anonymous signal channels."""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

HEADER_ROWS = 4
DATE_COL = 1
FIRST_CHANNEL_COL = 2


def find_raw_file(raw_dir):
    """Return the generic example workbook from the data directory."""
    files = sorted(Path(raw_dir).glob("synthetic_channel_data.xlsx"))
    if not files:
        raise FileNotFoundError(f"No synthetic_channel_data.xlsx found in {raw_dir}")
    return files[0]


def read_raw_sheet(path):
    """Read the workbook without consuming its four generic header rows."""
    return pd.read_excel(path, header=None)


def parse_header(raw):
    """Return channel metadata without source-system identifiers."""
    header = raw.iloc[:HEADER_ROWS, FIRST_CHANNEL_COL:]
    channels = header.iloc[2].astype("string").str.strip()
    if channels.isna().any() or channels.duplicated().any():
        raise ValueError("Every signal column must have a unique channel identifier.")
    if not channels.str.fullmatch(r"channel_\d{3}").all():
        raise ValueError("Channel identifiers must use the generic channel_NNN format.")

    records = []
    for position, channel in enumerate(channels, start=FIRST_CHANNEL_COL):
        records.append({
            "col_idx": position,
            "excel_col": get_column_letter(position + 1),
            "name": channel,
            "section": str(header.iloc[0, position - FIRST_CHANNEL_COL] or "group_00"),
            "description": str(header.iloc[3, position - FIRST_CHANNEL_COL] or "Signal"),
            "measurement": "signal",
            "unit": "",
        })
    return pd.DataFrame.from_records(records)


def to_numeric_clean(column):
    """Coerce a signal column to numbers and count invalid or boolean cells."""
    boolean = column.map(lambda value: isinstance(value, (bool, np.bool_)))
    text = column.map(lambda value: isinstance(value, str))
    numbers = pd.to_numeric(column.mask(boolean | text), errors="coerce")
    return numbers, int(text.sum()), int(boolean.sum())


def clean_values(raw, channel_dict):
    """Build a numeric, timestamp-indexed table and aggregate quality counts."""
    body = raw.iloc[HEADER_ROWS:].reset_index(drop=True)
    index = pd.to_datetime(body[DATE_COL], errors="coerce")
    if index.isna().any():
        raise ValueError(f"{int(index.isna().sum())} rows have an unreadable timestamp")

    columns = {}
    stats = []
    for _, row in channel_dict.iterrows():
        numbers, n_invalid, n_boolean = to_numeric_clean(body[row["col_idx"]])
        columns[row["name"]] = numbers.to_numpy()
        valid = numbers.dropna()
        stats.append({
            "n_rows": len(numbers),
            "n_valid": int(valid.size),
            "n_invalid": n_invalid,
            "n_boolean": n_boolean,
            "n_missing": int(numbers.isna().sum()),
            "pct_valid": round(100 * valid.size / len(numbers), 1),
            "n_unique": int(valid.nunique()),
        })

    values = pd.DataFrame(columns, index=pd.DatetimeIndex(index, name="timestamp"))
    channel_dict = pd.concat(
        [channel_dict.reset_index(drop=True), pd.DataFrame(stats)], axis=1
    )
    return values, channel_dict


def add_notes(channel_dict):
    """Add generic data-quality notes without attaching a physical meaning."""
    notes = []
    for _, row in channel_dict.iterrows():
        if row["n_valid"] == 0:
            notes.append("No numeric observations")
        elif row["pct_valid"] < 80:
            notes.append("High missingness")
        elif row["n_unique"] <= 3:
            notes.append("Low variation; review before modeling")
        else:
            notes.append("")
    channel_dict["notes"] = notes
    return channel_dict


def build_channel_dictionary(raw):
    """Build one anonymous metadata row per signal column."""
    return parse_header(raw)


def load_workbook(path):
    """Return a numeric signal table and its generic channel dictionary."""
    raw = read_raw_sheet(path)
    channel_dict = build_channel_dictionary(raw)
    values, channel_dict = clean_values(raw, channel_dict)
    channel_dict = add_notes(channel_dict)
    columns = [
        "name", "section", "description", "measurement", "unit", "excel_col",
        "pct_valid", "n_valid", "n_invalid", "n_boolean", "n_missing",
        "n_unique", "notes", "col_idx",
    ]
    return values, channel_dict[columns]


def write_markdown_dictionary(channel_dict, path):
    """Write a schema-only dictionary with no raw values or source identifiers."""
    lines = [
        "# Anonymous channel dictionary",
        "",
        "This schema lists generic channel identifiers and broad synthetic groups.",
        "It intentionally contains no source channel identifier, units, measurement ranges, or raw values.",
        "",
        "| Channel | Group | Description |",
        "|---|---|---|",
    ]
    for _, row in channel_dict.iterrows():
        lines.append(f"| {row['name']} | {row['section']} | Signal |")
    Path(path).write_text("\n".join(lines) + "\n", encoding="utf-8")
