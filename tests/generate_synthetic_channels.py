"""Create a fully synthetic workbook with anonymous channel identifiers."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.channel_io import HEADER_ROWS, load_workbook

CHANNEL_COUNT = 108
CORE_CHANNELS = [
    "channel_098", "channel_102", "channel_103", "channel_105",
    "channel_106", "channel_107", "channel_101", "channel_108",
    "channel_100", "channel_104", "channel_095", "channel_094",
    "channel_093", "channel_099",
]
SHUTDOWNS = [
    ("2033-03-18", "2033-03-20"),
    ("2033-06-19", "2033-06-25"),
    ("2033-08-10", "2033-08-15"),
]


def make_workbook(root: Path, seed: int = 5) -> Path:
    """Write two years of invented data; no external files are read."""
    dates = pd.date_range("2032-01-01", "2033-12-31", freq="D")
    rng = np.random.default_rng(seed)
    n_rows = len(dates)
    channel_names = [f"channel_{i:03d}" for i in range(1, CHANNEL_COUNT + 1)]
    values = {
        channel: rng.normal(10 + i / 20, 0.5 + (i % 5) / 10, n_rows)
        for i, channel in enumerate(channel_names, start=1)
    }

    year_two = dates.year == 2033
    progress = np.linspace(0, 1, int(year_two.sum()))
    values["channel_098"][year_two] = 0.2 + 0.12 * progress + rng.normal(0, 0.01, len(progress))
    values["channel_098"][dates >= "2033-12-20"] = 0.08
    values["channel_099"][year_two] *= 1 + 0.2 * progress
    values["channel_099"][dates >= "2033-11-01"] *= 2
    values["channel_102"][:] = 20 + rng.normal(0, 0.5, n_rows)
    values["channel_103"][:] = 18 + rng.normal(0, 0.5, n_rows)

    shutdown = pd.Series(False, index=dates)
    for start, end in SHUTDOWNS:
        shutdown.loc[start:end] = True
    values["channel_102"][shutdown.to_numpy()] = 1
    values["channel_103"][shutdown.to_numpy()] = 1

    for channel in channel_names:
        missing = rng.choice(n_rows, size=8, replace=False)
        values[channel][missing] = np.nan

    output = root / "data" / "raw" / "synthetic_channel_data.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic Data"

    for offset, channel in enumerate(channel_names, start=2):
        channel_number = offset - 1
        group = "group_01" if channel_number <= 7 else "group_02" if channel_number <= 72 else "group_03" if channel_number <= 94 else "group_04"
        sheet.cell(row=1, column=offset + 1, value=group)
        sheet.cell(row=2, column=offset + 1, value="generic signal")
        sheet.cell(row=3, column=offset + 1, value=channel)
        sheet.cell(row=4, column=offset + 1, value=f"Signal {offset - 1:03d}")
    sheet.cell(row=4, column=2, value="timestamp")

    for row, timestamp in enumerate(dates, start=HEADER_ROWS + 1):
        sheet.cell(row=row, column=2, value=timestamp.to_pydatetime())
        for column, channel in enumerate(channel_names, start=3):
            value = values[channel][row - HEADER_ROWS - 1]
            if np.isfinite(value):
                sheet.cell(row=row, column=column, value=float(value))

    workbook.save(output)
    return output


def verify(path: Path) -> None:
    raw = pd.read_excel(path, header=None)
    assert raw.shape == (735, CHANNEL_COUNT + 2), raw.shape
    values, channels = load_workbook(path)
    assert values.shape == (731, CHANNEL_COUNT), values.shape
    assert values.columns.tolist() == [f"channel_{i:03d}" for i in range(1, CHANNEL_COUNT + 1)]
    assert channels["name"].is_unique
    assert set(CORE_CHANNELS).issubset(values.columns)
    for start, end in SHUTDOWNS:
        assert (values.loc[start:end, "channel_102"] < 5).all()
        assert (values.loc[start:end, "channel_103"] < 5).all()
    assert values.loc["2033-11-15", "channel_099"] > values.loc["2033-03-01", "channel_099"]


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=".", help="Project root")
    parser.add_argument("--seed", type=int, default=5)
    arguments = parser.parse_args()
    destination = make_workbook(Path(arguments.root), seed=arguments.seed)
    verify(destination)
    print(f"Synthetic workbook verified: {destination}")
